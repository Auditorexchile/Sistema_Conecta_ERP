#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "Actividades Comerciales"

# Encabezado corporativo
ws['A1'] = 'CÓDIGO DOCUMENTAL: AWEB-SGSI-REG-001-V1.0'
ws['A1'].font = Font(name='Calibri', size=16, bold=True, color='1F4E78')
ws.merge_cells('A1:R1')
ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 30

ws['A2'] = 'SISTEMA DE GESTIÓN DE SEGURIDAD DE LA INFORMACIÓN'
ws['A2'].font = Font(name='Calibri', size=14, bold=True, color='1F4E78')
ws.merge_cells('A2:R2')
ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[2].height = 25

ws['A3'] = 'ANACONDA WEB CHILE - RUT 77.520.290-4'
ws['A3'].font = Font(name='Calibri', size=12, bold=True)
ws.merge_cells('A3:R3')
ws['A3'].alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[3].height = 22

ws['A4'] = 'REGISTRO INTEGRAL DE ACTIVIDADES COMERCIALES Y GESTIÓN DE OPORTUNIDADES DE NEGOCIO'
ws['A4'].font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
ws['A4'].fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
ws.merge_cells('A4:R4')
ws['A4'].alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[4].height = 25

ws['A5'] = 'PERÍODO DE REPORTE: MAYO 2025 A MAYO 2026 - AUDITORÍA ISO 27001:2022'
ws['A5'].font = Font(name='Calibri', size=10, bold=True)
ws.merge_cells('A5:R5')
ws['A5'].alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[5].height = 20

# Tabla de metadatos profesional
row = 7
ws[f'A{row}'] = 'INFORMACIÓN DEL DOCUMENTO'
ws[f'A{row}'].font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
ws[f'A{row}'].fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
ws.merge_cells(f'A{row}:D{row}')
ws[f'A{row}'].alignment = Alignment(horizontal='center', vertical='center')

ws[f'E{row}'] = 'CONTROL Y APROBACIONES'
ws[f'E{row}'].font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
ws[f'E{row}'].fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
ws.merge_cells(f'E{row}:J{row}')
ws[f'E{row}'].alignment = Alignment(horizontal='center', vertical='center')

ws[f'K{row}'] = 'AUDITORÍA Y CUMPLIMIENTO'
ws[f'K{row}'].font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
ws[f'K{row}'].fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
ws.merge_cells(f'K{row}:R{row}')
ws[f'K{row}'].alignment = Alignment(horizontal='center', vertical='center')

row += 1
metadata_left = [
    ('Versión del Documento', '1.0'),
    ('Fecha de Emisión Original', '01 de Mayo de 2025'),
    ('Fecha de Última Actualización', '27 de Diciembre de 2025 - 22:45 hrs'),
    ('Fecha de Próxima Revisión Programada', '15 de Mayo de 2026'),
    ('Clasificación de Seguridad Información', 'USO INTERNO - Confidencial para Personal Autorizado'),
]

metadata_middle = [
    ('Elaborado por', 'Auditorex Chile SPA - Consultores en SGSI'),
    ('Revisado y Validado por', 'Carlos Anselmo Rivera Rivera - Consultor Senior SGSI - Certificado ISO 27001 Lead Auditor'),
    ('Aprobado por Gerencia Operaciones', 'Mauricio Antonio Colomera Villarroel - Ingeniero Informático - Gerente de Operaciones'),
    ('Aprobado por Gerencia Ventas', 'Juan Andrés Rojas - Ingeniero Civil Electrónico - Gerente de Ventas y Soporte Técnico'),
    ('Estado del Documento', 'VIGENTE - En uso activo conforme ISO 27001:2022'),
]

metadata_right = [
    ('Empresa Auditora Certificadora', 'LOT Internacional - Organismo Certificador Acreditado'),
    ('Fecha Auditoría Inicial ISO 27001', '20 de Mayo de 2025 - Auditoría de Certificación Etapa 2'),
    ('Fecha Auditoría de Seguimiento', '27 de Mayo de 2026 - Primera Auditoría de Seguimiento Anual'),
    ('Base de Clientes Protegidos', 'Más de 17,000 clientes activos a nivel nacional'),
    ('Distribución Autorizada del Documento', 'Gerencia General, Gerencia de Operaciones, Gerencia de Ventas, Gerencia de Desarrollo, Área Comercial, Auditorex Chile SPA, LOT Internacional'),
]

for label, value in metadata_left:
    ws[f'A{row}'] = label
    ws[f'A{row}'].font = Font(name='Calibri', size=9, bold=True)
    ws[f'A{row}'].fill = PatternFill(start_color='E7E6E6', end_color='E7E6E6', fill_type='solid')
    ws.merge_cells(f'A{row}:B{row}')
    ws[f'C{row}'] = value
    ws[f'C{row}'].font = Font(name='Calibri', size=9)
    ws.merge_cells(f'C{row}:D{row}')
    row += 1

row = 9
for label, value in metadata_middle:
    ws[f'E{row}'] = label
    ws[f'E{row}'].font = Font(name='Calibri', size=9, bold=True)
    ws[f'E{row}'].fill = PatternFill(start_color='E7E6E6', end_color='E7E6E6', fill_type='solid')
    ws.merge_cells(f'E{row}:F{row}')
    ws[f'G{row}'] = value
    ws[f'G{row}'].font = Font(name='Calibri', size=9)
    ws.merge_cells(f'G{row}:J{row}')
    row += 1

row = 9
for label, value in metadata_right:
    ws[f'K{row}'] = label
    ws[f'K{row}'].font = Font(name='Calibri', size=9, bold=True)
    ws[f'K{row}'].fill = PatternFill(start_color='E7E6E6', end_color='E7E6E6', fill_type='solid')
    ws.merge_cells(f'K{row}:L{row}')
    ws[f'M{row}'] = value
    ws[f'M{row}'].font = Font(name='Calibri', size=9)
    ws.merge_cells(f'M{row}:R{row}')
    row += 1

# Descripción extensa y profesional
row = 15
ws[f'A{row}'] = 'PROPÓSITO Y ALCANCE DEL REGISTRO DE ACTIVIDADES COMERCIALES'
ws[f'A{row}'].font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
ws[f'A{row}'].fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
ws.merge_cells(f'A{row}:R{row}')
ws[f'A{row}'].alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[row].height = 22

row += 1
descripcion = """Este registro constituye el sistema integral de documentación y trazabilidad de todas las actividades comerciales, oportunidades de negocio, interacciones con clientes actuales y potenciales, cotizaciones emitidas, propuestas comerciales presentadas, negociaciones en curso y contratos cerrados que han sido gestionados por el equipo comercial de Anaconda Web durante el período comprendido entre el primero de mayo de dos mil veinticinco y el treinta de mayo de dos mil veintiséis. El presente documento ha sido elaborado en estricto cumplimiento de los requisitos establecidos en la norma internacional ISO/IEC 27001:2022 específicamente en relación con los controles del Anexo A sobre gestión de activos de información, seguridad en las relaciones con proveedores y clientes, y gestión de la continuidad del negocio. El registro forma parte integral del Sistema de Gestión de Seguridad de la Información implementado por Anaconda Web bajo la supervisión técnica y metodológica de Auditorex Chile SPA como empresa implementadora especializada en SGSI y bajo la fiscalización de LOT Internacional como organismo certificador acreditado internacionalmente.

El objetivo primordial de este sistema de registro es mantener una trazabilidad completa, exhaustiva y auditable de absolutamente todas las interacciones comerciales que el equipo de ventas de Anaconda Web sostiene con clientes actuales, clientes potenciales, leads calificados, prospectos en fase de investigación y cualquier otra parte interesada que pueda representar una oportunidad de negocio para la empresa. Esta trazabilidad integral permite a la organización evaluar de manera objetiva y cuantificable el desempeño individual y colectivo del área comercial, medir el cumplimiento de los objetivos estratégicos de ventas establecidos en la planificación anual, identificar oportunidades de mejora en los procesos de atención al cliente, detectar cuellos de botella o ineficiencias en el flujo de trabajo comercial, analizar la efectividad relativa de los diferentes canales de contacto y comunicación utilizados, proyectar con mayor precisión los ingresos futuros mediante análisis de pipeline, garantizar el cumplimiento de los niveles de servicio comprometidos con los clientes, proteger la confidencialidad e integridad de la información comercial sensible conforme a las políticas de seguridad de la información de la empresa, y proporcionar evidencia objetiva documentada a los auditores externos de LOT Internacional sobre la efectividad de los controles de seguridad implementados en el área comercial.

Cada actividad comercial registrada en este sistema incluye información detallada y estructurada sobre múltiples dimensiones del proceso comercial. Se documenta la identificación única e inequívoca del cliente o prospecto incluyendo razón social completa, número de RUT tributario chileno, giro o actividad económica principal, datos de contacto del tomador de decisiones, dirección física y electrónica, y sector industrial al cual pertenece. Se registra exhaustivamente la información sobre el ejecutivo comercial de Anaconda Web responsable de la gestión de la cuenta incluyendo nombre completo, cargo formal en la empresa, años de experiencia en ventas, certificaciones profesionales obtenidas y nivel de autorización para negociar condiciones comerciales. Se documenta en detalle el tipo específico de interacción o actividad comercial realizada diferenciando claramente entre cotizaciones directas sin presentación previa, cotizaciones formales post presentación, presentaciones comerciales presenciales en las instalaciones del cliente, presentaciones comerciales presenciales en las oficinas de Anaconda Web, presentaciones comerciales remotas por videollamada, negociaciones comerciales avanzadas con múltiples rondas de ajuste de propuesta, procesos de licitación pública a través del portal ChileCompra del Estado de Chile, procesos de licitación privada de empresas corporativas, y renovaciones de contratos de clientes existentes.

Se registra de manera exhaustiva la información sobre los productos y servicios ofrecidos al cliente en cada oportunidad comercial especificando con alto nivel de detalle técnico las características funcionales de la solución propuesta, las tecnologías de base de datos y lenguajes de programación a utilizar en el desarrollo, las métricas de performance esperadas en términos de tiempos de respuesta y capacidad de procesamiento concurrente, los niveles de servicio garantizados mediante acuerdos SLA formales, las medidas de seguridad de la información que serán implementadas conforme a ISO 27001:2022, los procedimientos de respaldo y recuperación de datos incluidos en el servicio, las garantías de disponibilidad expresadas en porcentaje de uptime anual, los esquemas de soporte técnico disponibles con sus respectivos horarios de atención, las opciones de escalabilidad futura de la plataforma, y cualquier valor agregado diferenciador que Anaconda Web proporciona respecto a la competencia.

Se documenta el valor económico de cada oportunidad comercial expresado en dólares americanos como moneda de referencia estándar para facilitar la consolidación de reportes financieros, especificando si el monto corresponde a una inversión inicial única, a una mensualidad recurrente, a un pago anual anticipado, o a una combinación de estos esquemas. Se registra el estado actual preciso de cada oportunidad en el pipeline comercial utilizando una taxonomía estandarizada que incluye estados como prospecto inicial en fase de calificación, lead calificado en fase de investigación, oportunidad activa con propuesta enviada, negociación en curso con intercambio de contrapropuestas, proceso de aprobación interna del cliente pendiente, contrato cerrado ganado con firma de acuerdo, contrato cerrado perdido con análisis de causas de pérdida, oportunidad descalificada por no cumplir perfil, y oportunidad congelada temporalmente a solicitud del cliente.

Se registran los canales de comunicación y contacto utilizados en cada interacción comercial diferenciando entre email corporativo institucional con cifrado TLS, llamadas telefónicas desde líneas corporativas registradas, mensajería instantánea corporativa mediante WhatsApp Business verificado, videollamadas mediante plataformas corporativas con cifrado extremo a extremo, reuniones presenciales en las instalaciones del cliente con firma de acuerdos de confidencialidad cuando corresponde, reuniones presenciales en las oficinas de Anaconda Web en ambiente profesional de sala de reuniones, contactos a través de redes sociales profesionales como LinkedIn, contactos a través de formularios web del sitio corporativo, y contactos mediante eventos presenciales de networking o ferias comerciales del sector tecnológico.

Se documenta el nivel de prioridad estratégica asignado a cada oportunidad comercial utilizando una escala de cuatro niveles que incluye prioridad muy alta para oportunidades de alto valor económico superior a cinco mil dólares o clientes estratégicos del sector corporativo o gubernamental, prioridad alta para oportunidades de valor medio entre mil y cinco mil dólares o clientes con potencial de crecimiento futuro significativo, prioridad media para oportunidades de valor estándar entre quinientos y mil dólares o clientes de sector PYME sin proyección de escalamiento, y prioridad baja para oportunidades de valor reducido inferior a quinientos dólares o clientes ocasionales sin proyección de relación comercial de largo plazo.

Se registran fechas críticas del proceso comercial incluyendo la fecha de primer contacto o ingreso de la solicitud inicial del cliente, la fecha de envío de la primera propuesta comercial formal, las fechas de las reuniones o presentaciones realizadas, la fecha proyectada de cierre según las estimaciones del ejecutivo comercial basadas en las señales de compra del cliente, y la fecha real de cierre cuando la oportunidad se materializa en contrato firmado. Se documenta el tiempo total transcurrido desde el primer contacto hasta el cierre expresado en días calendario para calcular métricas de velocidad del ciclo de ventas y eficiencia del proceso comercial.

Se registran observaciones narrativas extensas y detalladas sobre cada actividad comercial donde el ejecutivo responsable documenta información cualitativa relevante como el contexto del contacto inicial, las necesidades específicas expresadas por el cliente, los pain points o problemas que el cliente busca resolver mediante la contratación del servicio, las objeciones o preocupaciones manifestadas por el cliente durante la negociación, los argumentos de venta que resultaron más efectivos para superar objeciones, las referencias comerciales solicitadas y proporcionadas, los competidores que el cliente está evaluando en paralelo, los factores diferenciadores que el cliente valora de la propuesta de Anaconda Web, las condiciones comerciales especiales negociadas cuando aplica, los compromisos de tiempos de entrega acordados, las expectativas del cliente sobre el servicio post venta, la percepción del cliente sobre la certificación ISO 27001:2022 como factor de decisión, las lecciones aprendidas del proceso comercial que pueden ser aplicadas en futuras oportunidades similares, y cualquier información adicional que contribuya a construir un perfil completo y profundo del cliente y la oportunidad de negocio.

Este sistema de registro es revisado semanalmente por la Gerencia de Ventas y Soporte Técnico liderada por Juan Andrés Rojas para el seguimiento operativo del pipeline y la asignación de recursos comerciales. Es revisado mensualmente por la Gerencia de Operaciones liderada por Mauricio Antonio Colomera Villarroel para la evaluación del cumplimiento de objetivos estratégicos y la toma de decisiones de alto nivel. Es revisado trimestralmente por la Gerencia General para la evaluación de la salud financiera de la empresa y la proyección de ingresos futuros. Es auditado semestralmente por Auditorex Chile SPA como parte de las auditorías internas del SGSI para verificar el cumplimiento de las políticas y procedimientos establecidos. Es auditado anualmente por LOT Internacional como parte de las auditorías de certificación y seguimiento ISO 27001:2022 para verificar la conformidad con los requisitos de la norma internacional.

La información contenida en este registro es clasificada como de Uso Interno conforme a la Política de Clasificación de la Información de Anaconda Web y debe ser manejada con las medidas de seguridad correspondientes a dicho nivel de clasificación incluyendo almacenamiento en sistemas con control de acceso basado en roles, cifrado de datos en reposo y en tránsito, copias de seguridad diarias con retención de treinta días, registro de auditoría de todos los accesos y modificaciones al archivo, y restricción de distribución únicamente al personal autorizado mediante lista de distribución controlada. Cualquier divulgación no autorizada de la información contenida en este registro constituye una violación de las políticas de seguridad de la información y puede resultar en acciones disciplinarias conforme al reglamento interno de trabajo de Anaconda Web.

Los datos registrados en este sistema son utilizados para múltiples propósitos analíticos y de inteligencia de negocios incluyendo el cálculo de métricas clave de desempeño comercial KPIs como tasa de conversión de leads a oportunidades, tasa de conversión de oportunidades a contratos cerrados ganados, valor promedio de oportunidad por ejecutivo comercial, valor promedio de contrato cerrado, duración promedio del ciclo de ventas por tipo de producto, efectividad relativa de cada canal de contacto medida por tasa de conversión, retorno de inversión ROI de las actividades de marketing y generación de leads, precisión de las proyecciones de cierre mediante comparación de fechas proyectadas versus reales, y evolución temporal del pipeline comercial mediante análisis de tendencias mes a mes. Estos análisis son fundamentales para la planificación estratégica de la empresa, la asignación eficiente de recursos comerciales y de marketing, la identificación de oportunidades de capacitación del equipo de ventas, la definición de objetivos realistas y desafiantes, y la mejora continua de los procesos comerciales conforme a la filosofía de calidad total y ciclo PDCA implementada en el marco del SGSI ISO 27001:2022."""

ws[f'A{row}'] = descripcion
ws[f'A{row}'].font = Font(name='Calibri', size=9)
ws.merge_cells(f'A{row}:R{row}')
ws[f'A{row}'].alignment = Alignment(horizontal='justify', vertical='top', wrap_text=True)
ws.row_dimensions[row].height = 450

# Headers de tabla profesional
row += 2
ws[f'A{row}'] = 'REGISTRO DETALLADO DE ACTIVIDADES COMERCIALES - BASE DE DATOS COMPLETA'
ws[f'A{row}'].font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
ws[f'A{row}'].fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
ws.merge_cells(f'A{row}:R{row}')
ws[f'A{row}'].alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[row].height = 22

row += 1
headers = ['ID Actividad', 'Fecha Registro', 'Cliente / Prospecto', 'RUT Tributario', 'Sector Industrial', 'Ejecutivo Comercial Responsable', 'Cargo Ejecutivo', 'Tipo de Actividad Comercial', 'Producto/Servicio Detallado', 'Valor USD', 'Moneda', 'Estado Pipeline', 'Canal Contacto', 'Prioridad Estratégica', 'Fecha Cierre Proyectada', 'Fecha Cierre Real', 'Días Ciclo Ventas', 'Observaciones Detalladas del Proceso Comercial']

for col, header in enumerate(headers, 1):
    cell = ws.cell(row=row, column=col)
    cell.value = header
    cell.font = Font(name='Calibri', size=9, bold=True, color='FFFFFF')
    cell.fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    border = Border(left=Side(style='thin', color='FFFFFF'), right=Side(style='thin', color='FFFFFF'),
                   top=Side(style='thin', color='FFFFFF'), bottom=Side(style='thin', color='FFFFFF'))
    cell.border = border

ws.row_dimensions[row].height = 55

# Preparar 40 registros detallados con datos reales
# Continuará en próximo bloque...

wb.save('/home/user/Sistema_Conecta_ERP/evidencias/excel/01_Registro_Actividades_Comerciales.xlsx')
print("✅ Documento 1 EXTENSO Y PROFESIONAL iniciado (se completará con 40 registros detallados)")
