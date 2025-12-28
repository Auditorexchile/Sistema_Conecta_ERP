#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.drawing.image import Image as OpenpyxlImage
import os

wb = Workbook()

# ==================== PESTAÑA 1: PORTADA PROFESIONAL ====================
ws_portada = wb.active
ws_portada.title = "Portada"

# Configurar anchos de columna para la portada
for col in range(1, 11):
    ws_portada.column_dimensions[chr(64 + col)].width = 15

# Encabezado principal
ws_portada['A1'] = 'ANACONDA WEB CHILE'
ws_portada['A1'].font = Font(name='Calibri', size=36, bold=True, color='1F4E78')
ws_portada.merge_cells('A1:J1')
ws_portada['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws_portada.row_dimensions[1].height = 60

ws_portada['A2'] = 'RUT: 77.520.290-4'
ws_portada['A2'].font = Font(name='Calibri', size=14, bold=True, color='1F4E78')
ws_portada.merge_cells('A2:J2')
ws_portada['A2'].alignment = Alignment(horizontal='center', vertical='center')
ws_portada.row_dimensions[2].height = 25

# Espacio
ws_portada.row_dimensions[3].height = 20

# Título del documento
ws_portada['A4'] = 'SISTEMA DE GESTIÓN DE SEGURIDAD DE LA INFORMACIÓN'
ws_portada['A4'].font = Font(name='Calibri', size=18, bold=True, color='FFFFFF')
ws_portada['A4'].fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
ws_portada.merge_cells('A4:J4')
ws_portada['A4'].alignment = Alignment(horizontal='center', vertical='center')
ws_portada.row_dimensions[4].height = 35

ws_portada['A5'] = 'ISO/IEC 27001:2022'
ws_portada['A5'].font = Font(name='Calibri', size=16, bold=True, color='FFFFFF')
ws_portada['A5'].fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
ws_portada.merge_cells('A5:J5')
ws_portada['A5'].alignment = Alignment(horizontal='center', vertical='center')
ws_portada.row_dimensions[5].height = 30

# Espacio
ws_portada.row_dimensions[6].height = 30

# Nombre del documento
ws_portada['A7'] = 'REGISTRO INTEGRAL DE'
ws_portada['A7'].font = Font(name='Calibri', size=20, bold=True, color='1F4E78')
ws_portada.merge_cells('A7:J7')
ws_portada['A7'].alignment = Alignment(horizontal='center', vertical='center')
ws_portada.row_dimensions[7].height = 30

ws_portada['A8'] = 'ACTIVIDADES COMERCIALES'
ws_portada['A8'].font = Font(name='Calibri', size=24, bold=True, color='1F4E78')
ws_portada.merge_cells('A8:J8')
ws_portada['A8'].alignment = Alignment(horizontal='center', vertical='center')
ws_portada.row_dimensions[8].height = 35

ws_portada['A9'] = 'Y GESTIÓN DE OPORTUNIDADES DE NEGOCIO'
ws_portada['A9'].font = Font(name='Calibri', size=16, bold=True, color='1F4E78')
ws_portada.merge_cells('A9:J9')
ws_portada['A9'].alignment = Alignment(horizontal='center', vertical='center')
ws_portada.row_dimensions[9].height = 30

# Espacio
ws_portada.row_dimensions[10].height = 30

# Información del documento
row = 11
info_documento = [
    ('Código Documental:', 'AWEB-SGSI-REG-001-V1.0'),
    ('Versión:', '1.0'),
    ('Fecha de Emisión:', '01 de Mayo de 2025'),
    ('Fecha de Actualización:', '27 de Diciembre de 2025'),
    ('Clasificación:', 'USO INTERNO - Confidencial'),
    ('Estado:', 'VIGENTE'),
]

for label, value in info_documento:
    ws_portada[f'C{row}'] = label
    ws_portada[f'C{row}'].font = Font(name='Calibri', size=11, bold=True)
    ws_portada[f'C{row}'].fill = PatternFill(start_color='E7E6E6', end_color='E7E6E6', fill_type='solid')
    ws_portada.merge_cells(f'C{row}:D{row}')
    ws_portada[f'C{row}'].alignment = Alignment(horizontal='right', vertical='center')

    ws_portada[f'E{row}'] = value
    ws_portada[f'E{row}'].font = Font(name='Calibri', size=11)
    ws_portada.merge_cells(f'E{row}:H{row}')
    ws_portada[f'E{row}'].alignment = Alignment(horizontal='left', vertical='center')
    ws_portada.row_dimensions[row].height = 22
    row += 1

# Espacio
row += 2

# Periodo y auditoría
ws_portada[f'A{row}'] = 'PERÍODO DE REPORTE'
ws_portada[f'A{row}'].font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
ws_portada[f'A{row}'].fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
ws_portada.merge_cells(f'A{row}:J{row}')
ws_portada[f'A{row}'].alignment = Alignment(horizontal='center', vertical='center')
row += 1

ws_portada[f'A{row}'] = 'Mayo 2025 - Mayo 2026'
ws_portada[f'A{row}'].font = Font(name='Calibri', size=14, bold=True, color='1F4E78')
ws_portada.merge_cells(f'A{row}:J{row}')
ws_portada[f'A{row}'].alignment = Alignment(horizontal='center', vertical='center')
ws_portada.row_dimensions[row].height = 25
row += 2

# Auditoría
ws_portada[f'A{row}'] = 'INFORMACIÓN DE AUDITORÍA'
ws_portada[f'A{row}'].font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
ws_portada[f'A{row}'].fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
ws_portada.merge_cells(f'A{row}:J{row}')
ws_portada[f'A{row}'].alignment = Alignment(horizontal='center', vertical='center')
row += 1

info_auditoria = [
    ('Empresa Certificadora:', 'LOT Internacional'),
    ('Auditoría Inicial:', '20 de Mayo de 2025'),
    ('Próxima Auditoría:', '27 de Mayo de 2026'),
    ('Base de Clientes:', 'Más de 17,000 clientes activos'),
]

for label, value in info_auditoria:
    ws_portada[f'C{row}'] = label
    ws_portada[f'C{row}'].font = Font(name='Calibri', size=11, bold=True)
    ws_portada.merge_cells(f'C{row}:E{row}')
    ws_portada[f'C{row}'].alignment = Alignment(horizontal='right', vertical='center')

    ws_portada[f'F{row}'] = value
    ws_portada[f'F{row}'].font = Font(name='Calibri', size=11)
    ws_portada.merge_cells(f'F{row}:H{row}')
    ws_portada[f'F{row}'].alignment = Alignment(horizontal='left', vertical='center')
    ws_portada.row_dimensions[row].height = 22
    row += 1

# Espacio
row += 2

# Responsables
ws_portada[f'A{row}'] = 'ELABORACIÓN Y APROBACIÓN'
ws_portada[f'A{row}'].font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
ws_portada[f'A{row}'].fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
ws_portada.merge_cells(f'A{row}:J{row}')
ws_portada[f'A{row}'].alignment = Alignment(horizontal='center', vertical='center')
row += 1

responsables = [
    ('Elaborado por:', 'Auditorex Chile SPA'),
    ('Revisado por:', 'Carlos Anselmo Rivera Rivera - Consultor Senior SGSI'),
    ('Aprobado por:', 'Mauricio Antonio Colomera Villarroel - Gerente de Operaciones'),
    ('', 'Juan Andrés Rojas - Gerente de Ventas y Soporte Técnico'),
]

for label, value in responsables:
    if label:
        ws_portada[f'C{row}'] = label
        ws_portada[f'C{row}'].font = Font(name='Calibri', size=10, bold=True)
        ws_portada.merge_cells(f'C{row}:D{row}')
        ws_portada[f'C{row}'].alignment = Alignment(horizontal='right', vertical='center')

    ws_portada[f'E{row}'] = value
    ws_portada[f'E{row}'].font = Font(name='Calibri', size=10)
    ws_portada.merge_cells(f'E{row}:H{row}')
    ws_portada[f'E{row}'].alignment = Alignment(horizontal='left', vertical='center')
    ws_portada.row_dimensions[row].height = 20
    row += 1

# Pie de página
row += 3
ws_portada[f'A{row}'] = 'Este documento es propiedad de Anaconda Web Chile y contiene información confidencial protegida por las políticas de seguridad de la información del SGSI ISO 27001:2022'
ws_portada[f'A{row}'].font = Font(name='Calibri', size=8, italic=True, color='666666')
ws_portada.merge_cells(f'A{row}:J{row}')
ws_portada[f'A{row}'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws_portada.row_dimensions[row].height = 30

# ==================== PESTAÑA 2: ÍNDICE ====================
ws_indice = wb.create_sheet("Índice")

ws_indice['A1'] = 'ÍNDICE DEL DOCUMENTO'
ws_indice['A1'].font = Font(name='Calibri', size=18, bold=True, color='FFFFFF')
ws_indice['A1'].fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
ws_indice.merge_cells('A1:C1')
ws_indice['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws_indice.row_dimensions[1].height = 35

row = 3
indice_items = [
    ('1.', 'Portada', 'Información general del documento'),
    ('2.', 'Índice', 'Contenido del documento'),
    ('3.', 'Metadata y Propósito', 'Información detallada y alcance del registro'),
    ('4.', 'Registro de Actividades', 'Base de datos completa de actividades comerciales'),
    ('5.', 'Estadísticas y Análisis', 'Métricas, KPIs y análisis de desempeño'),
    ('6.', 'Metodología', 'Procedimientos y controles ISO 27001:2022'),
]

for num, titulo, descripcion in indice_items:
    ws_indice[f'A{row}'] = num
    ws_indice[f'A{row}'].font = Font(name='Calibri', size=12, bold=True, color='1F4E78')
    ws_indice[f'A{row}'].alignment = Alignment(horizontal='center', vertical='center')

    ws_indice[f'B{row}'] = titulo
    ws_indice[f'B{row}'].font = Font(name='Calibri', size=11, bold=True)

    ws_indice[f'C{row}'] = descripcion
    ws_indice[f'C{row}'].font = Font(name='Calibri', size=10)

    ws_indice.row_dimensions[row].height = 25
    row += 1

ws_indice.column_dimensions['A'].width = 8
ws_indice.column_dimensions['B'].width = 30
ws_indice.column_dimensions['C'].width = 60

# Continuará con las demás pestañas...
print("✅ Portada e Índice creados. Continuando con pestañas de datos...")

wb.save('/home/user/Sistema_Conecta_ERP/evidencias/excel/01_Registro_Actividades_Comerciales.xlsx')
print("✅ Documento 1 con PORTADA PROFESIONAL y PESTAÑAS guardado")
